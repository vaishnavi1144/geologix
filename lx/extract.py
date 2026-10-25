import os, sys, re, collections, numpy as np
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pymupdf as fitz
from depthcal import calibrate, depth_at, y_at
from legend import find_legend, kind_of
from load import load_drawings
from canvas import Canvas, is_rect_fill
from locate import column_extents
from classify import build_index, colour_close, row_runs
from column import vec_descriptor
from split import split_by_stamps
from descriptions import description_anchors, agree
from prim import col

from raster import extract_raster_page

def _cos(a, b): return float(a @ b / ((np.linalg.norm(a) * np.linalg.norm(b)) + 1e-9))

def consolidate_pages(page_results, page_errors=None):
    if not page_results:
        return {}
    if len(page_results) == 1 and not page_errors:
        return page_results[0]

    def _page_key(pr):
        if pr.get("rows"):
            return pr["rows"][0]["depth_start"]
        return pr.get("painted_depth_top", 0)

    sorted_pages = sorted(page_results, key=_page_key)
    unit = sorted_pages[0].get("unit", "m")
    min_thresh = 0.49 if unit == "ft" else 0.149

    consolidated_rows = []
    for pr in sorted_pages:
        page_rows = pr.get("rows", [])
        for row in page_rows:
            r = dict(row)
            if not consolidated_rows:
                consolidated_rows.append(r)
                continue

            last = consolidated_rows[-1]
            last_end = last["depth_end"]

            # Overlapping interval entirely before or at last_end
            if r["depth_end"] <= last_end + 0.001:
                continue

            # Starts before last_end: trim overlapping portion
            if r["depth_start"] < last_end:
                r["depth_start"] = last_end

            # Margin gap where r starts after last_end: bridge gap
            if r["depth_start"] > last_end:
                if r["lithology"] == last["lithology"]:
                    # Same lithology continuing across page boundary: merge!
                    last["depth_end"] = r["depth_end"]
                    last["confidence"] = round(min(last.get("confidence", 1.0), r.get("confidence", 1.0)), 2)
                    continue
                else:
                    # Different lithology: bridge to maintain strict contiguity
                    r["depth_start"] = last_end

            # Check if trimmed interval is too short (< min_thresh)
            thickness = round(r["depth_end"] - r["depth_start"], 2)
            if thickness < min_thresh:
                if r["lithology"] == last["lithology"]:
                    last["depth_end"] = r["depth_end"]
                    continue

            # Merge consecutive rows with identical lithology
            if r["lithology"] == last["lithology"]:
                last["depth_end"] = r["depth_end"]
                last["confidence"] = round(min(last.get("confidence", 1.0), r.get("confidence", 1.0)), 2)
            else:
                consolidated_rows.append(r)

    # Re-enforce strict contiguity and rounding
    for i in range(len(consolidated_rows)):
        cur = consolidated_rows[i]
        cur["depth_start"] = round(cur["depth_start"], 2)
        cur["depth_end"] = round(cur["depth_end"], 2)
        if i > 0:
            cur["depth_start"] = consolidated_rows[i - 1]["depth_end"]
            if cur["depth_end"] <= cur["depth_start"]:
                cur["depth_end"] = round(cur["depth_start"] + 0.15, 2)

    # Markers consolidation and deduplication across page boundaries
    all_markers = []
    for pr in sorted_pages:
        all_markers.extend(pr.get("markers", []))
    all_markers.sort(key=lambda m: (m.get("depth_mid", 0), m.get("depth_start", 0)))
    deduped_markers = []
    for m in all_markers:
        if deduped_markers:
            prev_m = deduped_markers[-1]
            if prev_m.get("lithology") == m.get("lithology") and abs(prev_m.get("depth_mid", 0) - m.get("depth_mid", 0)) < 0.25:
                continue
        deduped_markers.append(m)

    painted_top = min(pr["painted_depth_top"] for pr in sorted_pages)
    painted_bot = max(pr["painted_depth_bottom"] for pr in sorted_pages)

    warnings = []
    for pr in sorted_pages:
        warnings.extend(pr.get("warnings", []))
    if page_errors:
        warnings.extend(page_errors)

    marker_mode = "marker" if any(pr.get("marker_mode") == "marker" for pr in sorted_pages) else "variable-width"
    all_bounds = sorted(set(b for pr in sorted_pages for b in pr.get("bed_bounds", [])))

    return dict(
        rows=consolidated_rows,
        markers=deduped_markers,
        marker_mode=marker_mode,
        marker_count=len(deduped_markers),
        unit=unit,
        warnings=warnings,
        page_errors=page_errors or [],
        painted_depth_top=round(painted_top, 2),
        painted_depth_bottom=round(painted_bot, 2),
        bed_bounds=all_bounds,
        anchors=[a for pr in sorted_pages for a in pr.get("anchors", [])],
        doc=sorted_pages[0].get("doc"),
        page=sorted_pages[0].get("page"),
        cal=sorted_pages[0].get("cal"),
        column=sorted_pages[0].get("column"),
        candidates=sorted_pages[0].get("candidates"),
        legend=sorted_pages[0].get("legend"),
        lith=sorted_pages[0].get("lith"),
    )

def extract_page(pdf, page_no=None, column=None, verbose=False, min_pt=2.0):
    doc = fitz.open(pdf)
    total_pages = len(doc)

    if page_no is not None:
        return _extract_single_page(doc, page_no, column=column, verbose=verbose, min_pt=min_pt)

    if total_pages == 1:
        p = doc[0]
        drs = load_drawings(p)
        has_vector_page = len(drs) >= 5
        if has_vector_page:
            try:
                return _extract_single_page(doc, 0, column=column, verbose=verbose, min_pt=min_pt)
            except Exception as exc:
                err_msg = str(exc)
                if any(k in err_msg.lower() for k in ('legend', 'depth axis', 'column', 'feet', 'scanned', 'vector drawings')):
                    msg = err_msg.split(':', 1)[1].strip() if ':' in err_msg else err_msg
                    raise ValueError(msg)
                raise
        try:
            return extract_raster_page(doc, page_no=0, verbose=verbose)
        except Exception:
            raise ValueError('No valid lithology data found in document.')

    # Multi-page extraction workflow
    shared_legend = None
    shared_lith = None
    shared_column = column
    shared_unit = None

    # Pre-scan pages for shared legend / lithology / column
    for p_idx in range(total_pages):
        p = doc[p_idx]
        words = p.get_text('words')
        drs = load_drawings(p)
        if len(drs) < 5 or not words:
            continue
        try:
            cal_probe = calibrate(words)
            top_y = cal_probe['top'] - 60
        except Exception:
            cal_probe = None
            top_y = 500
        from legend import JUNK, find_legend, kind_of
        leg_probe = [e for e in find_legend(drs, words, top_y, p.rect.width)
                     if re.search('[A-Za-z]{3}', e['name']) and not JUNK.search(e['name'])]
        lith_probe = [e for e in leg_probe if e['section'] == 'lithology' or (e['section'] is None and kind_of(e['name']) == 'lithology')]
        if len(lith_probe) >= 3:
            shared_legend = leg_probe
            shared_lith = lith_probe
            if cal_probe and not shared_column:
                cv_init = Canvas(drs, 0, 0, p.rect.width, p.rect.height, rx=2.0, ry=1.0)
                cands_init = column_extents(cv_init, cal_probe, words, drs=drs)
                if cands_init:
                    shared_column = (cands_init[0]['x0'], cands_init[0]['x1'])
            break

    page_results = []
    page_errors = []
    last_cal = None

    for p_idx in range(total_pages):
        p = doc[p_idx]
        words = p.get_text('words')
        drs = load_drawings(p)

        # Check if page is blank or empty
        rect_fills = [d for d in drs if d['type'] in ('f', 'fs') and d.get('fill')]
        if len(words) == 0 and len(rect_fills) == 0:
            continue

        try:
            fb_cal = None
            if last_cal is not None:
                prev_bot_d = last_cal.get('last', 0)
                fb_cal = dict(last_cal)
                fb_cal['b'] = prev_bot_d

            res = _extract_single_page(
                doc, p_idx,
                column=column or shared_column,
                verbose=verbose,
                min_pt=min_pt,
                shared_legend=shared_legend,
                shared_lith=shared_lith,
                shared_column=shared_column,
                shared_unit=shared_unit,
                fallback_cal=fb_cal
            )
            page_results.append(res)
            if res.get('cal'):
                last_cal = res['cal']
            if not shared_column and res.get('column'):
                shared_column = res['column']
            if not shared_unit and res.get('unit'):
                shared_unit = res['unit']
        except Exception as exc:
            page_errors.append(f"Page {p_idx + 1}: {exc}")

    if not page_results:
        try:
            return extract_raster_page(doc, page_no=None, verbose=verbose)
        except Exception:
            pass
        for e in page_errors:
            if any(k in e.lower() for k in ('legend', 'depth axis', 'column', 'feet', 'scanned', 'vector drawings')):
                raise ValueError(e)
        if page_errors:
            raise ValueError(page_errors[0])
        raise ValueError('No valid lithology data found in document.')

    if len(page_results) == 1 and not page_errors:
        return page_results[0]

    return consolidate_pages(page_results, page_errors)

def _extract_single_page(doc, page_no, column=None, verbose=False, min_pt=2.0,
                         shared_legend=None, shared_lith=None, shared_column=None,
                         shared_unit=None, fallback_cal=None):
    p = doc[page_no]
    words = p.get_text('words')
    drs = load_drawings(p)

    # 5. FAIL LOUDLY: check vector drawings (scanned PDF detection)
    if len(drs) < 5:
        raise ValueError("No vector drawings found; PDF appears to be a scanned or raster document.")

    # 5. FAIL LOUDLY: depth axis calibration and depth unit check
    try:
        cal = calibrate(words)
    except Exception as exc:
        if fallback_cal is not None:
            cal = fallback_cal
        else:
            raise ValueError(f"No depth axis found: {exc}") from exc

    step = cal.get('grid', abs(cal.get('last', 0) - cal.get('first', 0)) / max(1, cal.get('n', 1) - 1))
    max_d = max(cal.get('first', 0), cal.get('last', 0))
    is_step_5_10 = (abs(step - 5) < 0.6 or abs(step - 10) < 0.6)
    has_ft_label = any(
        w[4].strip().lower() in ('ft', '(ft)', 'feet', '(feet)', 'depth(ft)', 'depth(feet)')
        for w in words
        if abs((w[0] + w[2]) / 2 - (cal['x'][0] + cal['x'][1]) / 2) < 60 and w[1] < cal['top'] + 50
    )
    is_feet = bool(shared_unit == 'ft' or (max_d > 5000 and is_step_5_10) or (has_ft_label and max_d > 1000) or (has_ft_label and step >= 50))
    depth_unit = 'ft' if is_feet else 'm'

    from legend import JUNK
    leg = [e for e in find_legend(drs, words, cal['top'] - 60, p.rect.width)
           if re.search('[A-Za-z]{3}', e['name']) and not JUNK.search(e['name'])]
    lith = [e for e in leg if e['section'] == 'lithology' or (e['section'] is None and kind_of(e['name']) == 'lithology')]

    # 5. FAIL LOUDLY: legend yields fewer than 3 lithologies
    if len(lith) < 3:
        if shared_lith and len(shared_lith) >= 3:
            lith = shared_lith
            leg = shared_legend or leg
        else:
            raise ValueError(f"Legend yields only {len(lith)} lithology entries (fewer than 3 required).")

    # Full-page initial canvas to discover colored lithology column across the entire page
    cv_initial = Canvas(drs, 0, 0, p.rect.width, p.rect.height, rx=2.0, ry=1.0)
    cands = column_extents(cv_initial, cal, words, drs=drs)
    if column:
        xa, xb = column
    else:
        # 5. FAIL LOUDLY: no lithology column found
        if not cands:
            if shared_column:
                xa, xb = shared_column
            else:
                raise ValueError('No coloured lithology column found in document.')
        else:
            xa, xb = cands[0]['x0'], cands[0]['x1']

    cw = xb - xa
    header_area = cal['top'] - 60

    # 1. EXTENT FROM PAINT, NOT LABELS
    # Derive y_top / y_bot from min/max of non-white fills inside column (bounded to column, y0 below header)
    fills_for_extent = [
        d['eff'] for d in drs
        if d['type'] in ('f', 'fs') and d.get('fill') and is_rect_fill(d)
        and not all(v > 0.97 for v in col(d['fill']))
        and d['eff'].x0 >= xa - 5 and d['eff'].x1 <= xb + 5
        and 0.3 * cw <= d['eff'].width <= 1.25 * cw and d['eff'].y0 >= header_area
    ]
    if fills_for_extent:
        y_top = min(e.y0 for e in fills_for_extent)
        y_bot = max(e.y1 for e in fills_for_extent)
    else:
        y_top = cal['top']
        y_bot = cal['bottom']

    painted_d_top = depth_at(cal, y_top)
    painted_d_bot = depth_at(cal, y_bot)

    # Header-text fallback: TD / Total Depth warning
    warnings = []
    if is_feet:
        warnings.append(f"Depth units detected as feet (ft) (max depth {max_d:.0f}, step {step:.1f}).")
    header_words = [w for w in words if w[1] < cal['top'] + 100]
    td_matches = []
    full_hdr_text = ' '.join(w[4] for w in sorted(header_words, key=lambda w: (w[1], w[0])))
    for m in re.finditer(r'\b(?:TD|TOTAL\s*DEPTH)\b\D{0,15}?(\d{3,5}(?:\.\d+)?)', full_hdr_text, re.I):
        try:
            td_matches.append(float(m.group(1)))
        except ValueError:
            pass
    label_spacing = cal.get('grid', abs(cal.get('last', 0) - cal.get('first', 0)) / max(1, cal.get('n', 1) - 1))
    if td_matches:
        td_val = max(td_matches)
        if abs(td_val - painted_d_bot) > label_spacing:
            w_msg = f"TD in header ({td_val:.1f}) conflicts with painted end ({painted_d_bot:.1f}) by more than 1 label spacing ({label_spacing:.1f})"
            warnings.append(w_msg)
            if verbose: print("WARNING:", w_msg)

    # 2. BEDS VS MARKERS, ADAPTIVELY
    # Collect all non-white rect fills in column inside [y_top - 1, y_bot + 1]
    all_col_fills = []
    for s, d in enumerate(drs):
        if d['type'] in ('f', 'fs') and d.get('fill') and is_rect_fill(d):
            c = col(d['fill'])
            if c and not all(v > 0.97 for v in c):
                e = d['eff']
                if e.x0 >= xa - 5 and e.x1 <= xb + 5 and 0.25 * cw <= e.width <= 1.25 * cw and e.y1 >= y_top - 1 and e.y0 <= y_bot + 1:
                    all_col_fills.append((s, d, c, e))

    col_edges = []
    for s, d, c, e in all_col_fills:
        col_edges.append(e.y0)
        col_edges.append(e.y1)
    col_edges = np.array(sorted(set(round(y, 3) for y in col_edges)))

    def _snap_edge(y_val):
        if len(col_edges) == 0:
            return y_val
        idx = np.searchsorted(col_edges, y_val)
        cands = col_edges[max(0, idx - 1):min(len(col_edges), idx + 2)]
        dists = np.abs(cands - y_val)
        if len(dists) > 0 and np.min(dists) <= 0.5:
            return float(cands[np.argmin(dists)])
        return y_val

    # d) Group tiles with identical (colour, y0, y1) before measuring width
    tile_groups = collections.defaultdict(list)
    for s, d, c, e in all_col_fills:
        k = (c, round(e.y0, 2), round(e.y1, 2))
        tile_groups[k].append((s, d, e))

    grouped = []
    for (c, y0, y1), items in tile_groups.items():
        min_x0 = min(e.x0 for _, _, e in items)
        max_x1 = max(e.x1 for _, _, e in items)
        w = max_x1 - min_x0
        h = y1 - y0
        min_s = min(s for s, _, _ in items)
        grouped.append({'c': c, 'y0': y0, 'y1': y1, 'x0': min_x0, 'x1': max_x1,
                        'w': w, 'h': h, 'seq': min_s, 'items': items})

    # a) Compute modal full-width of lithology fills
    widths = [round(g['w'], 1) for g in grouped if g['w'] >= 0.4 * cw]
    full_w = collections.Counter(widths).most_common(1)[0][0] if widths else cw

    full_fills = [g for g in grouped if g['w'] >= 0.8 * full_w]
    narrow_fills = [g for g in grouped if g['w'] < 0.8 * full_w]
    size_counts = collections.Counter((round(g['w'], 1), round(g['h'], 1)) for g in narrow_fills)
    repeating_sizes = {k for k, v in size_counts.items() if v >= 3}

    # b) & c) Identify markers or variable-width mode
    marker_seqs = set()
    marker_objs = []
    if repeating_sizes:
        marker_mode = 'marker'
        for g in narrow_fills:
            sz = (round(g['w'], 1), round(g['h'], 1))
            if sz not in repeating_sizes: continue
            ym = (g['y0'] + g['y1']) / 2
            under = [f for f in full_fills if f['y0'] - 0.5 <= ym <= f['y1'] + 0.5 and f['seq'] < g['seq']]
            if under:
                for s, _, _ in g['items']: marker_seqs.add(s)
                marker_objs.append(g)
    else:
        marker_mode = 'variable-width'

    if verbose or True:
        print(f"Marker mode: {marker_mode}, markers found: {len(marker_objs)}")

    # Second-pass Canvas over column's x-range and y-range with ry=0.1
    filtered_drs = [d for s, d in enumerate(drs) if s not in marker_seqs]
    min_pw = max(8.0, 0.3 * cw) if marker_mode == 'marker' else 0.4
    cv = Canvas(filtered_drs, xa - 1, y_top, xb + 1, y_bot, rx=1.0, ry=0.1, min_paint_w=min_pw)

    if verbose: print('column', round(xa), round(xb), '| legend lithologies', len(lith))
    idf = build_index(lith)
    for e in lith:
        e['desc'] = vec_descriptor(e['drawings'], (e['rect'][2] - e['rect'][0]) * (e['rect'][3] - e['rect'][1]))
    vis = [(s, d) for s, d in cv.visible_prims(xa, xb)]
    vis.sort(key=lambda t: t[1]['eff'].y0)
    ys = np.array([(d['eff'].y0 + d['eff'].y1) / 2 for _, d in vis]) if vis else np.array([])

    runs = row_runs(cv, xa, xb, 0, cv.H)
    bed_bounds = sorted(set([round(depth_at(cal, _snap_edge(cv.y_of_row(r[0]))), 2) for r in runs] + [round(depth_at(cal, _snap_edge(cv.y_of_row(r[1]))), 2) for r in runs]))
    ca = int((xa - cv.x0) / cv.rx); cb = int(np.ceil((xb - cv.x0) / cv.rx))
    cw_px = max(1, cb - ca)
    out = []
    for r0, r1, cid in runs:
        y0, y1 = _snap_edge(cv.y_of_row(r0)), _snap_edge(cv.y_of_row(r1))
        c = cv.colours[cid]
        if c is None or all(v > 0.97 for v in c):
            out.append(dict(y0=y0, y1=y1, name=None, conf=1.0, how='gap', rgb=c, runner=None)); continue
        cc = [e for e in lith if colour_close(e['bg'], c)]
        pool = cc if cc else lith
        # evidence from visible pattern primitives in this run
        i0, i1 = np.searchsorted(ys, y0), np.searchsorted(ys, y1)
        seg = vis[i0:i1]

        # 4. HONEST CONFIDENCE for how='colour'
        if len(cc) == 1:
            blk = cv.cid[r0:r1, ca:cb]
            cov = (blk == cid).sum(axis=1) / cw_px
            frac_95 = float((cov >= 0.95).mean())
            dist = float(np.linalg.norm(np.array(c) - np.array(cc[0]['bg']))) if cc[0]['bg'] else 0.0
            col_factor = max(0.0, 1.0 - min(1.0, dist * 10.0))
            honest_conf = round(max(0.1, min(0.99, frac_95 * col_factor)), 2)
            out.append(dict(y0=y0, y1=y1, name=cc[0]['name'], conf=honest_conf, how='colour', rgb=c, runner=None)); continue

        if cc and len(cc) >= 2:
            parts = split_by_stamps(seg, cc, y0, y1)
            if parts:
                for ya, yb, nm, fr in parts:
                    out.append(dict(y0=ya, y1=yb, name=nm, conf=round(min(0.95, 0.4 + 0.55 * fr), 2), how='colour+stamp-split', rgb=c,
                                    runner=next((e['name'] for e in cc if e['name'] != nm), None)))
                continue

        obs = collections.Counter(d['h'] for _, d in seg)
        union = set().union(*[set(e['prims']) for e in pool]) if pool else set()
        common = set.intersection(*[set(e['prims']) for e in pool]) if pool else set()
        base_seen = sum(n for h, n in obs.items() if h in common)      # a full tile is visible
        dist_map = {e['name']: set(e['prims']) - common for e in pool}
        sc = {}
        for e in pool:
            pos = sum(n * idf.get(h, .2) for h, n in obs.items() if h in dist_map[e['name']])
            neg = sum(n * idf.get(h, .2) for h, n in obs.items() if h in union and h not in e['prims'])
            absent = 0.0
            if cc and dist_map[e['name']] and pos == 0 and base_seen >= 2:
                absent = 0.5 * min(base_seen, 6) / 6.0 * max(idf.values(), default=1)   # tile visible, its mark is not
            sc[e['name']] = pos - 0.7 * neg - absent
        evidence = sum(n * idf.get(h, .2) for h, n in obs.items() if h in union)
        if evidence > 0:
            ranked = sorted(sc.items(), key=lambda kv: (-round(kv[1], 6), len(dist_map[kv[0]]), len(next(e for e in pool if e['name'] == kv[0])['prims'])))
            name = ranked[0][0]
            margin = ranked[0][1] - (ranked[1][1] if len(ranked) > 1 else 0)
            conf = round(max(0.1, min(0.95, 0.5 + margin / (evidence + 1e-9) * 0.5)), 2)
            out.append(dict(y0=y0, y1=y1, name=name, conf=conf, how='colour+pattern' if cc else 'pattern', rgb=c,
                            runner=ranked[1][0] if len(ranked) > 1 else None, seg=seg)); continue

        if cc:
            simp = sorted(cc, key=lambda e: (len(set(e['prims']) - common), len(e['prims'])))
            out.append(dict(y0=y0, y1=y1, name=simp[0]['name'], conf=0.3, how='colour-ambiguous', rgb=c,
                            runner=' / '.join(e['name'] for e in simp[1:4]))); continue

        # descriptor fallback (scale-tolerant)
        if seg and pool:
            dsc = vec_descriptor([d for _, d in seg], (xb - xa) * (y1 - y0))
            sc_desc = sorted(((_cos(dsc, e['desc']), e['name']) for e in pool), reverse=True)
            out.append(dict(y0=y0, y1=y1, name=sc_desc[0][1], conf=round(max(0, sc_desc[0][0]) * 0.6, 2), how='descriptor', rgb=c,
                            runner=sc_desc[1][1] if len(sc_desc) > 1 else None, seg=seg)); continue

        sig = (tuple(round(v, 1) for v in c), tuple(sorted(collections.Counter(d['h'] for _, d in seg).most_common(3))) if seg else ())
        out.append(dict(y0=y0, y1=y1, name=None, conf=0.0, how='cluster', rgb=c, runner=None, sig=sig))

    # Absorb slivers: ensure no interval shorter than min thickness (0.15 m or 0.5 ft)
    min_depth_len = 0.5 if depth_unit == 'ft' else 0.15
    min_pt_len = max(min_pt, min_depth_len * abs(cal['pt_per_unit']))
    merged = []
    for s in out:
        if merged and (s['y1'] - s['y0'] < min_pt_len):
            merged[-1]['y1'] = s['y1']; continue
        if merged and merged[-1]['name'] == s['name']:
            merged[-1]['y1'] = s['y1']; merged[-1]['conf'] = min(merged[-1]['conf'], s['conf']); continue
        merged.append(dict(s))

    # Convert y to depths and enforce strict contiguity
    rows = []
    for s in merged:
        d0 = depth_at(cal, s['y0'])
        d1 = depth_at(cal, s['y1'])
        cur_start = round(min(d0, d1), 2)
        cur_end = round(max(d0, d1), 2)
        if rows:
            cur_start = rows[-1]['depth_end']
            if cur_end <= cur_start:
                cur_end = round(cur_start + 0.01, 2)
        rows.append(dict(
            depth_start=cur_start,
            depth_end=cur_end,
            lithology=s['name'],
            confidence=s['conf'],
            method=s['how'],
            runner_up=s.get('runner'),
            y0=s['y0'],
            y1=s['y1'],
            page_no=page_no + 1
        ))

    # 3. MARKERS OUTPUT (never put in main rows)
    markers = []
    for g in marker_objs:
        d0 = depth_at(cal, g['y0'])
        d1 = depth_at(cal, g['y1'])
        dm = depth_at(cal, (g['y0'] + g['y1']) / 2)
        cc = [e for e in lith if colour_close(e['bg'], g['c'])]
        m_name = cc[0]['name'] if cc else 'Unknown'
        markers.append(dict(
            depth_start=round(min(d0, d1), 2),
            depth_end=round(max(d0, d1), 2),
            depth_mid=round(dm, 2),
            lithology=m_name,
            confidence=0.9 if cc else 0.5,
            runner_up=cc[1]['name'] if len(cc) > 1 else None,
            y0=g['y0'],
            y1=g['y1'],
            rgb=g['c'],
            page_no=page_no + 1
        ))

    anchors = description_anchors(words, cal['top'] - 5, cal['bottom'] + 5)
    # name legend-less clusters by majority vote of description anchors that fall inside them
    votes = collections.defaultdict(collections.Counter)
    for s_, r_ in zip(merged, rows):
        if s_.get('how') == 'cluster':
            for a in anchors:
                if s_['y0'] - 1 <= a['y'] <= s_['y1'] + 1: votes[s_['sig']][a['names'][0].title()] += 1
    for s_, r_ in zip(merged, rows):
        if s_.get('how') == 'cluster':
            v = votes.get(s_['sig'])
            if v: r_['lithology'] = v.most_common(1)[0][0]; r_['method'] = 'description-vote'; r_['confidence'] = 0.5
            else: r_['lithology'] = f"Unknown {abs(hash(s_['sig'])) % 1000:03d}"; r_['method'] = 'unlabelled-cluster'
    for r_ in rows:
        hits = [a for a in anchors if r_['y0'] - 1 <= a['y'] <= r_['y1'] + 1]
        r_['desc_lithology'] = ' | '.join(sorted({a['raw'] for a in hits})) or None
        r_['desc_agrees'] = (any(agree(r_['lithology'], a['names']) for a in hits)) if hits else None

    return dict(
        rows=rows,
        markers=markers,
        marker_mode=marker_mode,
        marker_count=len(markers),
        unit=depth_unit,
        warnings=warnings,
        painted_depth_top=round(min(painted_d_top, painted_d_bot), 2),
        painted_depth_bottom=round(max(painted_d_top, painted_d_bot), 2),
        bed_bounds=bed_bounds,
        anchors=anchors,
        cal=cal,
        column=(xa, xb),
        candidates=cands,
        legend=leg,
        lith=lith,
        page=p,
        doc=doc
    )

def export_to_excel(result, target):
    """
    Export extraction result to Excel (.xlsx) with:
    - Sheet 1: 'Lithology Data'
    - Sheet 2: 'Trace Markers' with note that marker height is symbol size, not bed thickness.
    """
    import pandas as pd
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    rows = result.get('rows', [])
    markers = result.get('markers', [])
    unit = result.get('unit', 'm')

    df_rows = pd.DataFrame(rows)
    col_map = {
        'depth_start': f'Depth Start ({unit})',
        'depth_end': f'Depth End ({unit})',
        'lithology': 'Lithology',
        'confidence': 'Confidence',
        'method': 'Method',
        'runner_up': 'Runner Up',
    }
    cols_use = [c for c in col_map if c in df_rows.columns]
    if cols_use:
        df_rows = df_rows[cols_use].rename(columns=col_map)

    with pd.ExcelWriter(target, engine='openpyxl') as writer:
        df_rows.to_excel(writer, index=False, sheet_name='Lithology Data')
        ws1 = writer.sheets['Lithology Data']

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=11)
        border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        for col_cells in ws1.columns:
            max_len = 0
            col_letter = col_cells[0].column_letter
            for cell in col_cells:
                cell.border = border
                if cell.row == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = regular_font
                    cell.alignment = Alignment(vertical="center")
                val_str = str(cell.value or "")
                if len(val_str) > max_len: max_len = len(val_str)
            ws1.column_dimensions[col_letter].width = max(max_len + 4, 15)

        # 3. MARKERS OUTPUT: second sheet "Trace Markers" with note
        if markers:
            df_m = pd.DataFrame(markers)
            m_map = {
                'depth_start': f'Depth Start ({unit})',
                'depth_end': f'Depth End ({unit})',
                'depth_mid': f'Depth Mid ({unit})',
                'lithology': 'Lithology',
                'confidence': 'Confidence',
                'runner_up': 'Runner Up',
            }
            m_cols = [c for c in m_map if c in df_m.columns]
            if m_cols:
                df_m = df_m[m_cols].rename(columns=m_map)

            df_m.to_excel(writer, index=False, sheet_name='Trace Markers', startrow=2)
            ws2 = writer.sheets['Trace Markers']

            ws2.cell(row=1, column=1, value="Note: Marker height is symbol size, not bed thickness.")
            ws2.cell(row=1, column=1).font = Font(name="Calibri", size=11, italic=True, bold=True, color="334155")

            for col_cells in ws2.iter_cols(min_row=3, max_row=len(df_m) + 3, min_col=1, max_col=len(df_m.columns)):
                max_len = 0
                col_letter = col_cells[0].column_letter
                for cell in col_cells:
                    cell.border = border
                    if cell.row == 3:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        cell.font = regular_font
                        cell.alignment = Alignment(vertical="center")
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len: max_len = len(val_str)
                ws2.column_dimensions[col_letter].width = max(max_len + 4, 15)

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(f'Usage: python {os.path.basename(__file__)} <pdf_path>')
        raise SystemExit(1)
    r = extract_page(sys.argv[1], verbose=True)
    print(f"Extracted {len(r['rows'])} rows, {len(r['markers'])} markers (mode: {r['marker_mode']})")
    for x in r['rows'][:15]:
        print(x['depth_start'], x['depth_end'], x['lithology'], x['confidence'], x['method'], x['runner_up'])

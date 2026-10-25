import gc
import io
import os
import re
import subprocess
import sys
import tempfile
import time
from pathlib import Path

from flask import Flask, jsonify, request, send_file, send_from_directory
from flask.json.provider import DefaultJSONProvider
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from lx.extract import extract_page, export_to_excel

FRONTEND_DIR = PROJECT_ROOT / "frontend"

class CustomJSONProvider(DefaultJSONProvider):
    def default(self, obj):
        if isinstance(obj, (np.floating, float)):
            return float(obj)
        if isinstance(obj, (np.integer, int)):
            return int(obj)
        if isinstance(obj, (np.bool_, bool)):
            return bool(obj)
        if isinstance(obj, np.ndarray):
            return obj.tolist()
        return super().default(obj)

app = Flask(__name__, static_folder=str(FRONTEND_DIR), static_url_path="")
app.json_provider_class = CustomJSONProvider
app.json = CustomJSONProvider(app)

LATEST_EXTRACTION = {
    "rows": [],
    "markers": [],
    "unit": "m",
    "filename": "lithology_extracted.xlsx",
}


@app.route("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


@app.route("/api/extract", methods=["POST"])
def api_extract():
    uploaded = request.files.get("file")
    if uploaded is None or uploaded.filename == "":
        return jsonify({"error": "Please upload a PDF file."}), 400

    filename = uploaded.filename.lower()
    if not filename.endswith(".pdf"):
        return jsonify({"error": "Only PDF files are supported."}), 400

    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            uploaded.save(tmp.name)
            temp_path = tmp.name

        result = extract_page(temp_path, verbose=False)
        rows = result.get("rows", [])
        markers = result.get("markers", [])
        unit = result.get("unit", "m")

        cleaned_rows = []
        for row in rows:
            d_start = row.get("depth_start")
            d_end = row.get("depth_end")
            conf = row.get("confidence")
            cleaned_rows.append({
                "depth_start": float(d_start) if d_start is not None else None,
                "depth_end": float(d_end) if d_end is not None else None,
                "lithology": str(row.get("lithology") or ""),
                "confidence": float(conf) if conf is not None else 1.0,
                "method": str(row.get("method") or ""),
                "runner_up": str(row.get("runner_up") or ""),
            })

        base_stem = Path(uploaded.filename).stem or "lithology"
        safe_base = re.sub(r'[<>:"/\\|?*]', '_', base_stem)
        LATEST_EXTRACTION["rows"] = cleaned_rows
        LATEST_EXTRACTION["markers"] = markers
        LATEST_EXTRACTION["unit"] = unit
        LATEST_EXTRACTION["filename"] = f"{safe_base}_extracted.xlsx"

        response = jsonify({
            "row_count": len(cleaned_rows),
            "rows": cleaned_rows,
            "unit": unit,
            "marker_count": len(markers),
            "markers": markers,
            "marker_mode": result.get("marker_mode", "variable-width"),
            "warnings": result.get("warnings", []),
            "filename": LATEST_EXTRACTION["filename"],
        })
        del result
        gc.collect()
        return response
    except Exception as exc:
        return jsonify({"error": str(exc)}), 500
    finally:
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except PermissionError:
                gc.collect()
                try:
                    os.remove(temp_path)
                except OSError:
                    pass


def generate_excel_response(rows, markers=None, unit="m", filename="lithology_extracted.xlsx", show_in_folder=True):
    df = pd.DataFrame(rows)
    column_mapping = {
        "depth_start": f"Depth Start ({unit})",
        "depth_end": f"Depth End ({unit})",
        "lithology": "Lithology",
        "confidence": "Confidence",
        "method": "Method",
        "runner_up": "Runner Up",
    }
    cols_to_use = [col for col in column_mapping.keys() if col in df.columns]
    if cols_to_use:
        df = df[cols_to_use].rename(columns=column_mapping)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        sheet_name = "Lithology Data"
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        worksheet = writer.sheets[sheet_name]

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=11)
        border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        for col in worksheet.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                cell.border = border
                if cell.row == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = regular_font
                    cell.alignment = Alignment(vertical="center")
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 15)

        # 3. MARKERS OUTPUT: export to second sheet "Trace Markers" with note
        if markers:
            markers_sheet_name = "Trace Markers"
            df_m = pd.DataFrame(markers)
            marker_mapping = {
                "depth_start": f"Depth Start ({unit})",
                "depth_end": f"Depth End ({unit})",
                "depth_mid": f"Depth Mid ({unit})",
                "lithology": "Lithology",
                "confidence": "Confidence",
                "runner_up": "Runner Up",
            }
            cols_to_use_m = [col for col in marker_mapping.keys() if col in df_m.columns]
            if cols_to_use_m:
                df_m = df_m[cols_to_use_m].rename(columns=marker_mapping)

            df_m.to_excel(writer, index=False, sheet_name=markers_sheet_name, startrow=2)
            ws_m = writer.sheets[markers_sheet_name]

            ws_m.cell(row=1, column=1, value="Note: Marker height is symbol size, not bed thickness.")
            ws_m.cell(row=1, column=1).font = Font(name="Calibri", size=11, italic=True, bold=True, color="334155")

            for col in ws_m.iter_cols(min_row=3, max_row=len(df_m) + 3, min_col=1, max_col=len(df_m.columns)):
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    cell.border = border
                    if cell.row == 3:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        cell.font = regular_font
                        cell.alignment = Alignment(vertical="center")
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws_m.column_dimensions[col_letter].width = max(max_len + 4, 15)

    excel_bytes = output.getvalue()

    # Automatically save directly to PC Downloads directory
    downloads_dir = Path.home() / "Downloads"
    safe_filename = re.sub(r'[<>:"/\\|?*]', '_', filename)
    if not safe_filename.endswith(".xlsx"):
        safe_filename = f"{safe_filename}.xlsx"
    target_path = downloads_dir / safe_filename

    try:
        with open(target_path, "wb") as f:
            f.write(excel_bytes)
    except PermissionError:
        ts = time.strftime("%Y%m%d_%H%M%S")
        safe_filename = f"{target_path.stem}_{ts}.xlsx"
        target_path = downloads_dir / safe_filename
        with open(target_path, "wb") as f:
            f.write(excel_bytes)
    except Exception as e:
        print(f"Direct file write to Downloads skipped: {e}")

    # Reveal in PC Downloads folder via Windows Explorer
    if show_in_folder and target_path.exists():
        try:
            subprocess.Popen(f'explorer.exe /select,"{str(target_path)}"', shell=True)
        except Exception as e:
            print(f"Failed to launch explorer: {e}")

    output.seek(0)
    response = send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=safe_filename,
    )
    response.headers["X-Saved-Path"] = str(target_path)
    response.headers["Access-Control-Expose-Headers"] = "X-Saved-Path"
    return response


@app.route("/api/download-excel", methods=["GET"])
def api_download_excel():
    rows = LATEST_EXTRACTION.get("rows", [])
    markers = LATEST_EXTRACTION.get("markers", [])
    unit = LATEST_EXTRACTION.get("unit", "m")
    if not rows:
        return "No extracted lithology data available. Please extract a PDF first.", 400

    filename = request.args.get("filename") or LATEST_EXTRACTION.get("filename", "lithology_extracted.xlsx")
    return generate_excel_response(rows, markers=markers, unit=unit, filename=filename, show_in_folder=True)


@app.route("/api/export-excel", methods=["POST"])
def api_export_excel():
    data = request.get_json(silent=True) or {}
    rows = data.get("rows", []) or LATEST_EXTRACTION.get("rows", [])
    markers = data.get("markers")
    if markers is None:
        markers = LATEST_EXTRACTION.get("markers", [])
    if not rows:
        return jsonify({"error": "No data available to export."}), 400

    filename = data.get("filename") or LATEST_EXTRACTION.get("filename", "lithology_extracted.xlsx")
    unit = data.get("unit") or LATEST_EXTRACTION.get("unit", "m")
    LATEST_EXTRACTION["rows"] = rows
    LATEST_EXTRACTION["markers"] = markers
    LATEST_EXTRACTION["unit"] = unit
    LATEST_EXTRACTION["filename"] = filename
    show_in_folder = data.get("show_in_folder", True)
    return generate_excel_response(rows, markers=markers, unit=unit, filename=filename, show_in_folder=show_in_folder)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False, use_reloader=False)
